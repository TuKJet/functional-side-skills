from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from copy import copy
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from . import __version__
from .config import DEFAULT_RULES
from .engine import assess_record, extract_metric, finite_confidence
from .parser import parse_cell


class AttendanceAuditError(RuntimeError):
    code = "ATTENDANCE_AUDIT_ERROR"


class SchemaMismatchError(AttendanceAuditError):
    code = "SCHEMA_MISMATCH"


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _parse_period(value: Any) -> tuple[datetime, datetime]:
    match = re.search(
        r"(?P<start>\d{4}-\d{2}-\d{2})\s+至\s+(?P<end>\d{4}-\d{2}-\d{2})",
        str(value or ""),
    )
    if not match:
        raise SchemaMismatchError("A1 未包含可识别的统计起止日期。")
    return (
        datetime.strptime(match.group("start"), "%Y-%m-%d"),
        datetime.strptime(match.group("end"), "%Y-%m-%d"),
    )


def _parse_generated_at(value: Any) -> datetime:
    match = re.search(r"(?P<value>\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2})", str(value or ""))
    if not match:
        raise SchemaMismatchError("A2 未包含可识别的报表生成时间。")
    return datetime.strptime(match.group("value"), "%Y-%m-%d %H:%M")


def _select_sheet(workbook: Any, daily_start_column: int) -> Any:
    for worksheet in workbook.worksheets:
        value = worksheet.cell(row=3, column=daily_start_column).value
        if str(value or "").strip() == "考勤结果":
            return worksheet
    raise SchemaMismatchError("未找到第 3 行以“考勤结果”标记的逐日考勤区域。")


def _validate_summary_headers(worksheet: Any, schema: dict[str, Any]) -> list[int]:
    columns: list[int] = []
    for letter, expected_header in schema["summary_columns"].items():
        column = column_index_from_string(letter)
        actual_header = str(worksheet.cell(row=3, column=column).value or "").strip()
        if actual_header != expected_header:
            raise SchemaMismatchError(
                f"汇总列 {letter} 表头不匹配：期望“{expected_header}”，实际“{actual_header}”。"
            )
        columns.append(column)
    return columns


def _validate_attendance_days_header(worksheet: Any, schema: dict[str, Any]) -> int:
    column = column_index_from_string(schema["attendance_days_column"])
    expected_header = schema["attendance_days_header"]
    actual_header = str(worksheet.cell(row=3, column=column).value or "").strip()
    if actual_header != expected_header:
        raise SchemaMismatchError(
            f"出勤天数列 {schema['attendance_days_column']} 表头不匹配："
            f"期望“{expected_header}”，实际“{actual_header}”。"
        )
    return column


def _is_positive_summary(value: Any) -> bool:
    if value in (None, ""):
        return False
    if isinstance(value, (int, float)):
        return value > 0
    text = str(value).strip()
    try:
        return float(text) > 0
    except ValueError:
        match = re.search(r"\d+(?:\.\d+)?", text)
        return bool(match and float(match.group(0)) > 0)


def _fill_key(cell: Any) -> str:
    if not cell.fill.fill_type:
        return "none"
    color = cell.fill.fgColor
    color_type = str(color.type or "none")
    value = str(getattr(color, color_type, ""))
    if color_type == "rgb":
        value = value.upper()
    return f"{color_type}:{value}"


def _text_screen_reason(value: Any, rules: dict[str, Any]) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    if any(keyword in text for keyword in rules["screening"]["explicit_anomaly_keywords"]):
        return "显式异常文本"
    if text.startswith("未排班"):
        punch_match = re.search(r"\n\(([^)]*)\)\s*$", text)
        if punch_match and any(
            punch.strip() not in {"", "-"} for punch in punch_match.group(1).split(",")
        ):
            return "未排班打卡"
    return None


def _person_from_row(worksheet: Any, row: int) -> dict[str, Any]:
    return {
        "name": str(worksheet.cell(row=row, column=1).value),
        "department": worksheet.cell(row=row, column=3).value,
        "employee_id": worksheet.cell(row=row, column=4).value,
    }


def prepare_analysis(
    input_path: str | Path,
    work_dir: str | Path,
) -> dict[str, Any]:
    input_path = Path(input_path).resolve()
    work_dir = Path(work_dir).resolve()
    if input_path.suffix.lower() != ".xlsx":
        raise SchemaMismatchError("当前版本只支持 .xlsx。")
    if not input_path.exists():
        raise AttendanceAuditError("输入文件不存在。")

    rules = DEFAULT_RULES
    schema = rules["schema"]
    daily_start_column = column_index_from_string(schema["daily_start_column"])
    workbook = load_workbook(input_path, data_only=False)
    worksheet = _select_sheet(workbook, daily_start_column)
    period_start, period_end = _parse_period(worksheet[schema["period_cell"]].value)
    generated_at = _parse_generated_at(worksheet[schema["generated_at_cell"]].value)
    expected_days = (period_end.date() - period_start.date()).days + 1
    actual_days = worksheet.max_column - daily_start_column + 1
    if actual_days != expected_days:
        raise SchemaMismatchError(
            f"逐日列数量与统计区间不一致：期望 {expected_days} 列，实际 {actual_days} 列。"
        )
    summary_columns = _validate_summary_headers(worksheet, schema)
    attendance_days_column = _validate_attendance_days_header(worksheet, schema)

    data_start_row = int(schema["data_start_row"])
    people_rows: list[int] = []
    summary_candidate_rows: set[int] = set()
    colored_columns_by_row: dict[int, set[int]] = defaultdict(set)
    text_columns_by_row: dict[int, dict[int, str]] = defaultdict(dict)
    normal_fill_keys = set(rules["screening"]["normal_fill_keys"])
    for row in range(data_start_row, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=1).value
        if name in (None, ""):
            continue
        people_rows.append(row)
        if any(
            _is_positive_summary(worksheet.cell(row=row, column=column).value)
            for column in summary_columns
        ):
            summary_candidate_rows.add(row)
        for column in range(daily_start_column, worksheet.max_column + 1):
            cell = worksheet.cell(row=row, column=column)
            if cell.value not in (None, ""):
                if _fill_key(cell) not in normal_fill_keys:
                    colored_columns_by_row[row].add(column)
                text_reason = _text_screen_reason(cell.value, rules)
                if text_reason:
                    text_columns_by_row[row][column] = text_reason

    candidate_rows = sorted(
        summary_candidate_rows.union(colored_columns_by_row).union(text_columns_by_row)
    )
    context_radius = int(rules["screening"]["neighbor_days"])
    assessments: list[dict[str, Any]] = []
    review_items: list[dict[str, Any]] = []
    candidate_cells_examined = 0
    for row in candidate_rows:
        person = _person_from_row(worksheet, row)
        if row in summary_candidate_rows:
            target_columns = set(range(daily_start_column, worksheet.max_column + 1))
        else:
            target_columns = set(colored_columns_by_row[row]).union(text_columns_by_row[row])
        context_columns = set(target_columns)
        for column in target_columns:
            for offset in range(-context_radius, context_radius + 1):
                neighbor = column + offset
                if daily_start_column <= neighbor <= worksheet.max_column:
                    context_columns.add(neighbor)

        records: list[dict[str, Any]] = []
        for column in sorted(context_columns):
            cell = worksheet.cell(row=row, column=column)
            screening_reasons: list[str] = []
            if row in summary_candidate_rows:
                screening_reasons.append("T-Z汇总")
            if column in colored_columns_by_row[row]:
                screening_reasons.append("特殊颜色")
            if column in text_columns_by_row[row]:
                screening_reasons.append(text_columns_by_row[row][column])
            records.append(
                {
                    "source_cell": f"{worksheet.title}!{cell.coordinate}",
                    "row": row,
                    "column": column,
                    "date": period_start + timedelta(days=column - daily_start_column),
                    "person": person,
                    "parsed": parse_cell(cell.value, rules),
                    "screening_reasons": screening_reasons,
                }
            )
        record_indexes = {record["column"]: index for index, record in enumerate(records)}
        for column in sorted(target_columns):
            index = record_indexes[column]
            record = records[index]
            assessment = assess_record(record, records, index, rules, generated_at)
            candidate_cells_examined += 1
            if (
                assessment.findings
                or assessment.review_item
                or assessment.corrected_text != assessment.parsed.raw
            ):
                assessments.append(assessment.to_dict())
            if assessment.review_item:
                review_items.append(assessment.review_item)

    analysis_path = work_dir / "analysis.json"
    queue_path = work_dir / "agent_review_queue.json"
    review_dir = work_dir / "agent_review_items"
    if review_dir.exists():
        for stale_item in review_dir.glob("review_*.json"):
            stale_item.unlink()
    if review_items:
        review_dir.mkdir(parents=True, exist_ok=True)
    elif review_dir.exists() and not any(review_dir.iterdir()):
        review_dir.rmdir()
    queue_items = []
    for index, review_item in enumerate(review_items, start=1):
        item_path = review_dir / f"review_{index:05d}.json"
        _write_json(item_path, review_item)
        queue_items.append(
            {
                "review_id": review_item["review_id"],
                "review_type": review_item["review_type"],
                "item_file": str(item_path),
            }
        )
    analysis = {
        "schema_version": "2.0",
        "engine_version": __version__,
        "source_path": str(input_path),
        "source_sha256": sha256_file(input_path),
        "sheet_name": worksheet.title,
        "daily_start_column": daily_start_column,
        "attendance_days_column": attendance_days_column,
        "max_column": worksheet.max_column,
        "max_row": worksheet.max_row,
        "candidate_rows": candidate_rows,
        "assessments": assessments,
    }
    queue = {
        "schema_version": "2.0",
        "privacy_note": (
            "Read one item_file at a time. Items omit name, employee ID, department, position, "
            "and attendance group. Never read analysis.json into model context."
        ),
        "item_count": len(review_items),
        "items": queue_items,
    }
    _write_json(analysis_path, analysis)
    _write_json(queue_path, queue)
    return {
        "status": "needs_agent_review" if review_items else "prepared",
        "processed_people": len(people_rows),
        "candidate_people": len(candidate_rows),
        "summary_candidate_people": len(summary_candidate_rows),
        "processed_days": expected_days,
        "daily_cells_total": len(people_rows) * expected_days,
        "candidate_cells_examined": candidate_cells_examined,
        "notable_cells": len(assessments),
        "special_color_cells": sum(len(columns) for columns in colored_columns_by_row.values()),
        "explicit_text_cells": sum(len(columns) for columns in text_columns_by_row.values()),
        "agent_review_items": len(review_items),
        "analysis_file": str(analysis_path),
        "agent_review_queue": str(queue_path),
    }


def _load_decisions(path: str | Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    decision_path = Path(path)
    if not decision_path.exists():
        raise AttendanceAuditError("Agent 决策文件不存在。")
    payload = json.loads(decision_path.read_text(encoding="utf-8"))
    decisions = payload.get("decisions")
    if not isinstance(decisions, list):
        raise AttendanceAuditError("Agent 决策文件缺少 decisions 数组。")
    result: dict[str, dict[str, Any]] = {}
    for decision in decisions:
        review_id = decision.get("review_id")
        action = decision.get("action")
        if not review_id or action not in {"keep", "accept_suggestion", "replace"}:
            raise AttendanceAuditError("Agent 决策包含无效的 review_id 或 action。")
        if review_id in result:
            raise AttendanceAuditError("Agent 决策包含重复的 review_id。")
        confidence = finite_confidence(decision.get("confidence"), 0.0)
        if not 0 <= confidence <= 1:
            raise AttendanceAuditError("Agent 决策 confidence 必须位于 0 到 1。")
        result[review_id] = {**decision, "confidence": confidence}
    return result


def _punch_suffix(text: str) -> str:
    match = re.search(r"\n\([^)]*\)\s*$", text)
    return match.group(0) if match else ""


def _resolve_review_text(
    assessment: dict[str, Any],
    decisions: dict[str, dict[str, Any]],
) -> tuple[str, dict[str, Any] | None]:
    review = assessment.get("review_item")
    if not review:
        return assessment["corrected_text"], None
    decision = decisions.get(review["review_id"])
    if not decision or decision["action"] == "keep":
        return assessment["corrected_text"], decision
    if decision["action"] == "accept_suggestion":
        candidate = review.get("candidate") or {}
        corrected = candidate.get("proposed_text")
        if not corrected:
            raise AttendanceAuditError("该 review_id 没有可接受的 proposed_text。")
    else:
        corrected = decision.get("corrected_text")
        if not isinstance(corrected, str) or not corrected.strip():
            raise AttendanceAuditError("replace 决策必须提供 corrected_text。")
        corrected = corrected.strip()
    original_suffix = _punch_suffix(assessment["parsed"]["raw"])
    if original_suffix and _punch_suffix(corrected) != original_suffix:
        raise AttendanceAuditError("Agent 决策不得修改原始打卡括号内容。")
    return corrected, decision


def _comment_text(
    assessment: dict[str, Any],
    final_text: str,
    decision: dict[str, Any] | None,
) -> str:
    findings = assessment.get("findings", [])
    lines = [f"考勤审计版本：{__version__}"]
    lines.append("规则：" + ", ".join(item["rule_id"] for item in findings))
    lines.extend(f"- {item['message']}" for item in findings)
    lines.append(f"原值：{assessment['parsed']['raw']}")
    if final_text != assessment["parsed"]["raw"]:
        lines.append(f"矫正值：{final_text}")
    if decision:
        lines.append(f"Agent 决策：{decision.get('action')}")
        if decision.get("reason"):
            lines.append(f"Agent 说明：{decision['reason']}")
        lines.append(f"Agent 置信度：{decision.get('confidence', 0):.2f}")
    elif assessment.get("review_item"):
        lines.append("Agent 决策：未提供，保留原值并等待人工审查。")
    return "\n".join(lines)


def _append_comment(cell: Any, text: str) -> None:
    if cell.comment and cell.comment.text:
        text = f"{cell.comment.text}\n\n{text}"
    cell.comment = Comment(text, "Attendance Audit")


def _same_cell_style(source: Any, corrected: Any) -> bool:
    return all(
        (
            copy(source.font) == copy(corrected.font),
            copy(source.fill) == copy(corrected.fill),
            copy(source.border) == copy(corrected.border),
            copy(source.alignment) == copy(corrected.alignment),
            source.number_format == corrected.number_format,
            copy(source.protection) == copy(corrected.protection),
        )
    )


def _monthly_metrics(
    worksheet: Any, row: int, start_column: int, end_column: int, rules: dict[str, Any]
) -> dict[str, Any]:
    summary = {20: 0, 21: 0, 22: 0, 23: 0, 24: 0, 25: 0, 26: 0}
    missing_punch_days = 0
    unclassified_missing_cells: list[str] = []
    unmeasured_time_cells: list[str] = []
    for column in range(start_column, end_column + 1):
        parsed = parse_cell(worksheet.cell(row=row, column=column).value, rules)
        status = parsed.shift_status or ""
        missing_in = "上班缺卡" in status
        missing_out = "下班缺卡" in status
        if missing_in:
            summary[20] += 1
        if missing_out:
            summary[21] += 1
        if missing_in ^ missing_out:
            missing_punch_days += 1
        elif "缺卡" in status:
            unclassified_missing_cells.append(worksheet.cell(row=row, column=column).coordinate)
        if "迟到" in status:
            summary[22] += 1
            minutes = extract_metric(status, "迟到")
            summary[23] += minutes
            if minutes == 0 and not re.search(r"迟到0分钟", status):
                unmeasured_time_cells.append(worksheet.cell(row=row, column=column).coordinate)
        if "早退" in status:
            summary[24] += 1
            minutes = extract_metric(status, "早退")
            summary[25] += minutes
            if minutes == 0 and not re.search(r"早退0分钟", status):
                unmeasured_time_cells.append(worksheet.cell(row=row, column=column).coordinate)
        if "旷工" in status:
            summary[26] += 1
    return {
        "summary": summary,
        "missing_punch_days": missing_punch_days,
        "late_minutes": summary[23],
        "early_minutes": summary[25],
        "unclassified_missing_cells": sorted(set(unclassified_missing_cells)),
        "unmeasured_time_cells": sorted(set(unmeasured_time_cells)),
    }


def _decimal_value(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def _decimal_text(value: Decimal) -> str:
    return format(value, "f").rstrip("0").rstrip(".") or "0"


def _attendance_base_days(cell: Any) -> Decimal | None:
    if cell.comment and cell.comment.text:
        matches = re.findall(r"基准出勤天数：(-?\d+(?:\.\d+)?)", cell.comment.text)
        if matches:
            return _decimal_value(matches[-1])
    return _decimal_value(cell.value)


def _attendance_alignment(
    cell: Any,
    monthly: dict[str, Any],
    settings: dict[str, Any],
) -> dict[str, Any]:
    late_minutes = int(monthly["late_minutes"])
    early_minutes = int(monthly["early_minutes"])
    total_minutes = late_minutes + early_minutes
    increment = Decimal(str(settings["rounding_increment_days"]))
    workday_minutes = Decimal(str(settings["workday_minutes"]))
    missing_deduction = Decimal(str(settings["missing_punch_deduction_days"])) * int(
        monthly["missing_punch_days"]
    )
    time_deduction = Decimal("0")
    if total_minutes:
        raw_days = Decimal(total_minutes) / workday_minutes
        time_deduction = (raw_days / increment).to_integral_value(
            rounding=ROUND_CEILING
        ) * increment
    base_days = _attendance_base_days(cell)
    result = {
        "base_days": base_days,
        "missing_punch_days": int(monthly["missing_punch_days"]),
        "missing_deduction_days": missing_deduction,
        "late_minutes": late_minutes,
        "early_minutes": early_minutes,
        "time_deduction_days": time_deduction,
        "unclassified_missing_cells": monthly["unclassified_missing_cells"],
        "unmeasured_time_cells": monthly["unmeasured_time_cells"],
        "aligned_days": None,
    }
    if (
        base_days is None
        or monthly["unclassified_missing_cells"]
        or monthly["unmeasured_time_cells"]
    ):
        return result
    result["aligned_days"] = max(
        Decimal("0"),
        base_days - missing_deduction - time_deduction,
    )
    return result


def _markdown_cell(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def _write_shift_correction_report(
    output_path: Path,
    details: list[dict[str, Any]],
) -> None:
    direction_counts: dict[str, int] = defaultdict(int)
    for detail in details:
        direction_counts[f"{detail['original_shift']} → {detail['corrected_shift']}"] += 1
    lines = [
        "# A/B 班排班矫正汇总",
        "",
        f"- 矫正次数：{len(details)}",
        f"- 涉及人数：{len({detail['employee_id'] or detail['name'] for detail in details})}",
    ]
    for direction, count in sorted(direction_counts.items()):
        lines.append(f"- {direction}：{count}")
    lines.extend(
        [
            "",
            "| 序号 | 姓名 | 工号 | 部门 | 日期 | 原班次 | 矫正班次 | 源单元格 |",
            "|---:|---|---|---|---|---|---|---|",
        ]
    )
    for index, detail in enumerate(details, start=1):
        values = [
            index,
            detail["name"],
            detail["employee_id"],
            detail["department"],
            detail["date"],
            detail["original_shift"],
            detail["corrected_shift"],
            detail["source_cell"],
        ]
        lines.append("| " + " | ".join(_markdown_cell(value) for value in values) + " |")
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _weekday(date_value: str) -> str:
    parsed = datetime.strptime(date_value, "%Y-%m-%d").date()
    return "一二三四五六日"[parsed.weekday()]


def _status_line(value: Any) -> str:
    return str(value or "").splitlines()[0].strip()


def _day_handling(day: dict[str, Any]) -> str:
    original = _status_line(day.get("original_text"))
    corrected = _status_line(day.get("corrected_text"))
    if day.get("auto_corrected") and corrected and corrected != original:
        handling = f"已自动将“{original}”修正为“{corrected}”，原打卡记录不变。"
    elif day.get("needs_human_review"):
        handling = "保留原记录，未自动修改。"
    else:
        handling = "异常已确认并保留原记录。"
    if day.get("needs_human_review"):
        handling += f" 需人工确认：{'；'.join(day['messages'])}"
    return handling


def _summarize_daily_anomalies(
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    people: dict[str, dict[str, Any]] = {}
    for row in rows:
        date_value = row.get("date")
        if not date_value:
            continue
        identity = str(row.get("employee_id") or row.get("name") or row["source_cell"])
        person = people.setdefault(
            identity,
            {
                "name": row.get("name"),
                "employee_id": row.get("employee_id"),
                "days": {},
            },
        )
        day = person["days"].setdefault(
            date_value,
            {
                "categories": [],
                "messages": [],
                "original_text": row.get("original_text"),
                "corrected_text": row.get("corrected_text"),
                "auto_corrected": False,
                "needs_human_review": False,
            },
        )
        if row["category"] not in day["categories"]:
            day["categories"].append(row["category"])
        if row["message"] not in day["messages"]:
            day["messages"].append(row["message"])
        day["auto_corrected"] = day["auto_corrected"] or bool(row.get("auto_corrected"))
        day["needs_human_review"] = day["needs_human_review"] or bool(row.get("needs_human_review"))
        if row.get("corrected_text") != row.get("original_text"):
            day["corrected_text"] = row.get("corrected_text")

    person_summaries: list[dict[str, Any]] = []
    llm_items: list[dict[str, Any]] = []
    for person in people.values():
        date_lines: list[str] = []
        automatic_lines: list[str] = []
        review_lines: list[str] = []
        for date_value, day in sorted(person["days"].items()):
            label = f"{date_value[5:]}（{_weekday(date_value)}）"
            anomaly = "、".join(day["categories"])
            detail = "；".join(day["messages"])
            handling = _day_handling(day)
            date_lines.append(f"{label} {anomaly}：{detail}")
            if day["auto_corrected"]:
                automatic_lines.append(f"{label} {handling.split(' 需人工确认：', 1)[0]}")
            if day["needs_human_review"]:
                review_lines.append(f"{label} {detail}")
            llm_items.append(
                {
                    "name": person["name"],
                    "date": date_value,
                    "anomaly": anomaly,
                    "detail": detail,
                    "handling": handling,
                }
            )
        person_summaries.append(
            {
                "name": person["name"],
                "employee_id": person["employee_id"],
                "anomaly_days": len(person["days"]),
                "date_lines": date_lines,
                "automatic_lines": automatic_lines,
                "review_lines": review_lines,
            }
        )
    return person_summaries, llm_items


def _create_anomaly_workbook(rows: list[dict[str, Any]], output_path: Path) -> None:
    person_summaries, _ = _summarize_daily_anomalies(rows)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "异常审查"
    headers = [
        "序号",
        "姓名",
        "工号",
        "异常天数",
        "异常日期与类型",
        "自动处理结果",
        "待人工确认",
    ]
    worksheet.append(headers)
    for index, person in enumerate(person_summaries, start=1):
        worksheet.append(
            [
                index,
                person["name"],
                person["employee_id"],
                person["anomaly_days"],
                "\n".join(person["date_lines"]),
                "\n".join(person["automatic_lines"]) or "无",
                "\n".join(person["review_lines"]) or "无",
            ]
        )
        row_number = index + 1
        line_count = max(
            person["anomaly_days"],
            len(person["automatic_lines"]),
            len(person["review_lines"]),
            1,
        )
        worksheet.row_dimensions[row_number].height = min(180, max(30, line_count * 21))

    header_fill = PatternFill("solid", fgColor="9C0006")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9D9D9")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_gray)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[3].alignment = Alignment(horizontal="center", vertical="top")
        if row[5].value != "无":
            row[5].fill = PatternFill("solid", fgColor="E2F0D9")
        if row[6].value != "无":
            row[6].fill = PatternFill("solid", fgColor="FFF2CC")
    worksheet.row_dimensions[1].height = 26
    widths = {
        "A": 8,
        "B": 12,
        "C": 12,
        "D": 10,
        "E": 62,
        "F": 54,
        "G": 54,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(1, worksheet.max_row)}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def finalize_analysis(
    analysis_path: str | Path,
    output_dir: str | Path,
    decisions_path: str | Path | None = None,
) -> dict[str, Any]:
    analysis_path = Path(analysis_path).resolve()
    output_dir = Path(output_dir).resolve()
    analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
    source_path = Path(analysis["source_path"])
    if sha256_file(source_path) != analysis["source_sha256"]:
        raise AttendanceAuditError("源文件在准备与最终处理之间发生变化，已停止。")
    rules = DEFAULT_RULES
    decisions = _load_decisions(decisions_path)
    workbook = load_workbook(source_path, data_only=False)
    worksheet = workbook[analysis["sheet_name"]]

    anomaly_rows: list[dict[str, Any]] = []
    shift_correction_details: list[dict[str, Any]] = []
    final_by_cell: dict[str, str] = {}
    unresolved_rows: set[int] = set()
    for assessment in analysis["assessments"]:
        final_text, decision = _resolve_review_text(assessment, decisions)
        if assessment.get("review_item") and (not decision or decision.get("action") == "keep"):
            unresolved_rows.add(int(assessment["row"]))
        final_by_cell[assessment["source_cell"]] = final_text
        cell = worksheet.cell(row=assessment["row"], column=assessment["column"])
        if final_text != assessment["parsed"]["raw"]:
            cell.value = final_text
            final_parsed = parse_cell(final_text, rules)
            original_shift_id = assessment["parsed"].get("shift_id")
            if (
                original_shift_id in {"SHIFT_A", "SHIFT_B"}
                and final_parsed.shift_id in {"SHIFT_A", "SHIFT_B"}
                and final_parsed.shift_id != original_shift_id
            ):
                person = assessment["person"]
                shift_correction_details.append(
                    {
                        "name": person.get("name"),
                        "employee_id": person.get("employee_id"),
                        "department": person.get("department"),
                        "date": assessment["date"],
                        "source_cell": assessment["source_cell"],
                        "original_shift": assessment["parsed"].get("shift_label"),
                        "corrected_shift": final_parsed.shift_label,
                    }
                )
        if assessment["findings"]:
            _append_comment(cell, _comment_text(assessment, final_text, decision))
        for finding in assessment["findings"]:
            person = assessment["person"]
            anomaly_rows.append(
                {
                    "needs_human_review": finding["needs_human_review"],
                    "severity": finding["severity"],
                    "category": finding["category"],
                    "name": person.get("name"),
                    "employee_id": person.get("employee_id"),
                    "department": person.get("department"),
                    "date": assessment["date"],
                    "source_cell": assessment["source_cell"],
                    "original_text": assessment["parsed"]["raw"],
                    "corrected_text": final_text,
                    "auto_corrected": final_text != assessment["parsed"]["raw"],
                    "screening_reasons": assessment.get("screening_reasons", []),
                    "rule_id": finding["rule_id"],
                    "message": finding["message"],
                    "agent_action": decision.get("action") if decision else None,
                    "agent_reason": decision.get("reason") if decision else None,
                    "confidence": decision.get("confidence") if decision else finding["confidence"],
                }
            )

    summary_changes = 0
    attendance_days_changes = 0
    attendance_days_pending = 0
    attendance_days_column = int(analysis["attendance_days_column"])
    alignment_settings = rules["attendance_alignment"]
    person_rows = [
        int(row)
        for row in analysis.get(
            "candidate_rows",
            sorted({item["row"] for item in analysis["assessments"]}),
        )
    ]
    for row in person_rows:
        monthly = _monthly_metrics(
            worksheet,
            row,
            int(analysis["daily_start_column"]),
            int(analysis["max_column"]),
            rules,
        )
        if row not in unresolved_rows:
            for column, new_value in monthly["summary"].items():
                cell = worksheet.cell(row=row, column=column)
                old_value = cell.value or 0
                if old_value == new_value:
                    continue
                cell.value = new_value
                _append_comment(
                    cell,
                    (
                        f"考勤审计版本：{__version__}\n"
                        f"规则：SUMMARY_RECONCILIATION\n"
                        f"原值：{old_value}\n矫正值：{new_value}\n"
                        "依据：矫正后的逐日考勤结果重新汇总。"
                    ),
                )
                summary_changes += 1

        alignment_cell = worksheet.cell(row=row, column=attendance_days_column)
        alignment = _attendance_alignment(alignment_cell, monthly, alignment_settings)
        has_alignment_deduction = bool(
            alignment["missing_punch_days"]
            or alignment["late_minutes"]
            or alignment["early_minutes"]
            or alignment["unclassified_missing_cells"]
            or alignment["unmeasured_time_cells"]
        )
        if not has_alignment_deduction:
            continue

        alignment_blocked = row in unresolved_rows or alignment["aligned_days"] is None
        if alignment_blocked:
            attendance_days_pending += 1
            if row in unresolved_rows:
                message = "该人员仍有未决逐日记录，R 列保持原值，待人工确认后重新对齐。"
            elif alignment["unclassified_missing_cells"]:
                cells = "、".join(alignment["unclassified_missing_cells"])
                message = f"缺卡记录未明确为单侧缺卡（{cells}），R 列保持原值。"
            elif alignment["unmeasured_time_cells"]:
                cells = "、".join(alignment["unmeasured_time_cells"])
                message = f"迟到或早退记录缺少可计算分钟数（{cells}），R 列保持原值。"
            else:
                message = "R 列原出勤天数不是可计算数值，已保持原值。"
            _append_comment(
                alignment_cell,
                (
                    f"考勤审计版本：{__version__}\n"
                    "规则：ATTENDANCE_DAYS_ALIGNMENT_PENDING\n"
                    f"{message}"
                ),
            )
            continue

        base_days = alignment["base_days"]
        aligned_days = alignment["aligned_days"]
        current_days = _decimal_value(alignment_cell.value)
        value_changed = current_days != aligned_days
        if value_changed:
            alignment_cell.value = float(aligned_days)
            _append_comment(
                alignment_cell,
                (
                    f"考勤审计版本：{__version__}\n"
                    "规则：ATTENDANCE_DAYS_ALIGNMENT\n"
                    f"基准出勤天数：{_decimal_text(base_days)}\n"
                    f"缺卡日期：{alignment['missing_punch_days']} 天，"
                    f"扣减：{_decimal_text(alignment['missing_deduction_days'])} 天\n"
                    f"迟到：{alignment['late_minutes']} 分钟；"
                    f"早退：{alignment['early_minutes']} 分钟\n"
                    f"迟到早退扣减：{_decimal_text(alignment['time_deduction_days'])} 天\n"
                    f"对齐后出勤天数：{_decimal_text(aligned_days)}"
                ),
            )
            attendance_days_changes += 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    corrected_path = output_dir / f"{source_path.stem}_考勤数据对齐_{__version__}_{timestamp}.xlsx"
    anomaly_path = output_dir / f"{source_path.stem}_考勤异常_{__version__}_{timestamp}.xlsx"
    llm_summary_path = output_dir / f"{source_path.stem}_异常摘要_{__version__}_{timestamp}.json"
    person_summaries, llm_items = _summarize_daily_anomalies(anomaly_rows)
    shift_report_path: Path | None = None
    if len(shift_correction_details) > int(alignment_settings["shift_report_threshold"]):
        shift_report_path = (
            output_dir / f"{source_path.stem}_班次修复汇总_{__version__}_{timestamp}.md"
        )
        _write_shift_correction_report(shift_report_path, shift_correction_details)
    _append_comment(
        worksheet["A2"],
        (
            f"考勤审计版本：{__version__}\n"
            f"处理时间：{datetime.now().isoformat(timespec='seconds')}\n"
            f"日级异常：{len(llm_items)} 项，涉及 {len(person_summaries)} 人\n"
            f"月度汇总矫正：{summary_changes}\n"
            f"出勤天数对齐：{attendance_days_changes}\n"
            f"A/B 班矫正：{len(shift_correction_details)}"
        ),
    )
    workbook.save(corrected_path)
    _create_anomaly_workbook(anomaly_rows, anomaly_path)
    llm_summary_path.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "source_workbook": source_path.name,
                "items": llm_items,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    source_check = load_workbook(source_path, read_only=False, data_only=False)
    corrected_check = load_workbook(corrected_path, read_only=False, data_only=False)
    anomaly_check = load_workbook(anomaly_path, read_only=True, data_only=True)
    if corrected_check.sheetnames != source_check.sheetnames:
        raise AttendanceAuditError("矫正副本的工作表结构发生意外变化。")
    source_sheet = source_check[analysis["sheet_name"]]
    corrected_sheet = corrected_check[analysis["sheet_name"]]
    if set(source_sheet.merged_cells.ranges) != set(corrected_sheet.merged_cells.ranges):
        raise AttendanceAuditError("矫正副本的合并单元格发生意外变化。")
    for row in range(1, 5):
        for column in range(1, int(analysis["max_column"]) + 1):
            if (
                corrected_sheet.cell(row=row, column=column).value
                != source_sheet.cell(row=row, column=column).value
            ):
                raise AttendanceAuditError("矫正副本的固定表头验证失败。")
    for row in range(1, int(analysis["max_row"]) + 1):
        if source_sheet.row_dimensions[row].height != corrected_sheet.row_dimensions[row].height:
            raise AttendanceAuditError("矫正副本的行高发生意外变化。")
        for column in range(1, int(analysis["max_column"]) + 1):
            if not _same_cell_style(
                source_sheet.cell(row=row, column=column),
                corrected_sheet.cell(row=row, column=column),
            ):
                coordinate = f"{get_column_letter(column)}{row}"
                raise AttendanceAuditError(f"矫正副本的单元格样式发生意外变化：{coordinate}。")
    for column in range(1, int(analysis["max_column"]) + 1):
        letter = get_column_letter(column)
        if (
            source_sheet.column_dimensions[letter].width
            != corrected_sheet.column_dimensions[letter].width
        ):
            raise AttendanceAuditError("矫正副本的列宽发生意外变化。")
    if "异常审查" not in anomaly_check.sheetnames:
        raise AttendanceAuditError("异常审查工作簿验证失败。")
    anomaly_sheet = anomaly_check["异常审查"]
    if anomaly_sheet.max_row != len(person_summaries) + 1:
        raise AttendanceAuditError("异常审查工作簿未保持一人一行。")
    if any(not item.get("date") for item in llm_items):
        raise AttendanceAuditError("LLM 异常摘要包含无日期记录。")
    if sha256_file(source_path) != analysis["source_sha256"]:
        raise AttendanceAuditError("源文件在最终处理过程中发生变化。")

    return {
        "status": "completed",
        "engine_version": __version__,
        "source_unchanged": True,
        "corrected_cells": sum(
            1
            for assessment in analysis["assessments"]
            if final_by_cell[assessment["source_cell"]] != assessment["parsed"]["raw"]
        ),
        "summary_cells_corrected": summary_changes,
        "attendance_days_cells_corrected": attendance_days_changes,
        "attendance_days_alignment_pending": attendance_days_pending,
        "anomaly_rows": len(person_summaries),
        "daily_anomaly_items": len(llm_items),
        "shift_corrections": len(shift_correction_details),
        "shift_correction_people": len(
            {detail["employee_id"] or detail["name"] for detail in shift_correction_details}
        ),
        "shift_correction_report": str(shift_report_path) if shift_report_path else None,
        "agent_decisions_applied": sum(
            1 for decision in decisions.values() if decision.get("action") != "keep"
        ),
        "corrected_workbook": str(corrected_path),
        "anomaly_workbook": str(anomaly_path),
        "llm_summary_file": str(llm_summary_path),
    }
