import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from attendance_audit.workbook import finalize_analysis, prepare_analysis, sha256_file


def _build_workbook(path: Path, unknown: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "月度汇总"
    sheet["A1"] = "月度汇总 统计日期：2026-07-01 至 2026-07-03"
    sheet["A2"] = "报表生成时间：2026-07-04 10:00"
    headers = ["姓名", "考勤组", "部门", "工号", "职位", "UserId"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index).value = header
    sheet["T3"] = "上班缺卡次数"
    sheet["U3"] = "下班缺卡次数"
    sheet["V3"] = "迟到次数"
    sheet["W3"] = "迟到时长"
    sheet["X3"] = "早退次数"
    sheet["Y3"] = "早退时长"
    sheet["Z3"] = "旷工天数"
    sheet["R3"] = "出勤天数"
    sheet["AG3"] = "考勤结果"
    sheet["AG4"] = "1"
    sheet["AH4"] = "2"
    sheet["AI4"] = "3"
    sheet["A5"] = "Sensitive Name"
    sheet["B5"] = "Sensitive Group"
    sheet["C5"] = "Sensitive Department"
    sheet["D5"] = "0001"
    sheet["E5"] = "Sensitive Title"
    sheet["F5"] = "Sensitive UserId"
    sheet["R5"] = 3
    for column in range(20, 27):
        sheet.cell(row=5, column=column).value = 0
    sheet["V5"] = 1
    sheet["W5"] = 10
    sheet["AG5"] = "研发A班:正常\n(09:00,18:00)"
    sheet["AH5"] = "研发A班:正常\n(09:20,18:26)"
    sheet["AI5"] = (
        "研发A班:正常,特殊备案待确认\n(09:00,18:00)" if unknown else "未排班\n(09:00,18:35)"
    )
    sheet["AG5"].fill = PatternFill("solid", fgColor="FFF2CC")
    workbook.save(path)


def test_end_to_end_corrects_copy_and_builds_anomaly_workbook(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    work_dir = tmp_path / "work"
    output_dir = tmp_path / "outputs"
    _build_workbook(source)
    source_hash = sha256_file(source)

    prepared = prepare_analysis(source, work_dir)
    assert prepared["agent_review_items"] == 0
    assert not (work_dir / "agent_review_items").exists()
    decisions = work_dir / "agent_decisions.json"
    decisions.write_text('{"schema_version":"2.0","decisions":[]}', encoding="utf-8")
    finalized = finalize_analysis(prepared["analysis_file"], output_dir, decisions)

    assert sha256_file(source) == source_hash
    corrected = load_workbook(finalized["corrected_workbook"])
    sheet = corrected["月度汇总"]
    assert sheet["AH5"].value == "研发A班:上班迟到10分钟\n(09:20,18:26)"
    assert sheet["AI5"].value == "未排班并打卡\n(09:00,18:35)"
    assert sheet["V5"].value == 1
    assert sheet["W5"].value == 10
    assert sheet["R5"].value == 2.9
    assert sheet["AG5"].fill.fgColor.rgb == "00FFF2CC"
    assert sheet["AH5"].comment is not None

    anomaly = load_workbook(finalized["anomaly_workbook"], read_only=True)
    rows = list(anomaly["异常审查"].iter_rows(values_only=True))
    headers = list(rows[0])
    assert headers == [
        "序号",
        "姓名",
        "工号",
        "异常天数",
        "异常日期与类型",
        "原始考勤数据",
        "处理结果",
        "待人工确认",
    ]
    assert len(rows) == 2
    assert rows[1][headers.index("异常天数")] == 2
    anomaly_detail = rows[1][headers.index("异常日期与类型")]
    assert "07-02（四） 迟到" in anomaly_detail
    assert "07-03（五） 未排班打卡" in anomaly_detail
    assert "出勤天数已对齐" not in anomaly_detail
    original_detail = rows[1][headers.index("原始考勤数据")]
    assert "07-02（四） 研发A班:正常 ｜ (09:20,18:26)" in original_detail
    assert "07-03（五） 未排班 ｜ (09:00,18:35)" in original_detail
    assert len(original_detail.splitlines()) == 2
    handling_detail = rows[1][headers.index("处理结果")]
    assert "07-02（四） 已自动将“研发A班:正常”修正为“研发A班:上班迟到10分钟”" in handling_detail
    assert "07-03（五） 已自动将“未排班”修正为“未排班并打卡”" in handling_detail
    assert len(handling_detail.splitlines()) == 2

    llm_summary = json.loads(Path(finalized["llm_summary_file"]).read_text(encoding="utf-8"))
    assert len(llm_summary["items"]) == 2
    assert all(item["name"] == "Sensitive Name" for item in llm_summary["items"])
    assert all(item["date"] for item in llm_summary["items"])
    assert {item["anomaly"] for item in llm_summary["items"]} == {
        "迟到",
        "未排班打卡",
    }

    anomaly_editable = load_workbook(finalized["anomaly_workbook"])
    anomaly_sheet = anomaly_editable["异常审查"]
    assert anomaly_sheet.auto_filter.ref == "A1:H2"
    assert not anomaly_sheet.tables
    assert anomaly_sheet["H2"].fill.fill_type == "solid"
    with ZipFile(finalized["anomaly_workbook"]) as archive:
        names = set(archive.namelist())
        assert "xl/tables/table1.xml" not in names
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert sheet_xml.count("<autoFilter") == 1
        assert "<tableParts" not in sheet_xml


def test_agent_queue_omits_person_and_organization_metadata(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    work_dir = tmp_path / "work"
    _build_workbook(source, unknown=True)

    prepared = prepare_analysis(source, work_dir)
    queue_text = Path(prepared["agent_review_queue"]).read_text(encoding="utf-8")
    queue = json.loads(queue_text)

    assert queue["item_count"] == 1
    assert "Sensitive Name" not in queue_text
    assert "Sensitive Group" not in queue_text
    assert "Sensitive Department" not in queue_text
    assert queue["items"][0]["review_type"] == "unsupported_text"
    item_text = Path(queue["items"][0]["item_file"]).read_text(encoding="utf-8")
    assert "Sensitive Name" not in item_text
    assert "Sensitive Group" not in item_text
    assert "Sensitive Department" not in item_text
    assert len(item_text) < 2000


def test_python_auto_corrects_decisive_shift_misclassification(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "input.xlsx"
    work_dir = tmp_path / "work"
    output_dir = tmp_path / "outputs"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    sheet["AG5"] = "研发B班:正常\n(11:50,21:05)"
    sheet["AH5"] = "研发A班:上班迟到50分钟\n(10:00,21:00)"
    sheet["AI5"] = "研发B班:正常\n(11:55,21:02)"
    workbook.save(source)

    prepared = prepare_analysis(source, work_dir)
    assert prepared["agent_review_items"] == 0

    from attendance_audit.config import DEFAULT_RULES

    monkeypatch.setitem(DEFAULT_RULES["attendance_alignment"], "shift_report_threshold", 0)
    finalized = finalize_analysis(prepared["analysis_file"], output_dir)
    corrected = load_workbook(finalized["corrected_workbook"])
    assert corrected["月度汇总"]["AH5"].value == "研发B班:正常\n(10:00,21:00)"
    assert corrected["月度汇总"]["V5"].value == 0
    assert corrected["月度汇总"]["W5"].value == 0
    assert corrected["月度汇总"]["R5"].value == 3
    assert finalized["shift_corrections"] == 1
    assert Path(finalized["shift_correction_report"]).exists()
    anomaly = load_workbook(finalized["anomaly_workbook"], read_only=True)
    values = list(anomaly["异常审查"].iter_rows(values_only=True))
    headers = list(values[0])
    assert len(values) == 2
    assert "班次错排已矫正" in values[1][headers.index("异常日期与类型")]
    assert "研发A班:上班迟到50分钟 ｜ (10:00,21:00)" in values[1][headers.index("原始考勤数据")]
    assert "已自动将" in values[1][headers.index("处理结果")]
    assert "无需人工确认" in values[1][headers.index("待人工确认")]
    anomaly_editable = load_workbook(finalized["anomaly_workbook"])
    assert anomaly_editable["异常审查"]["H2"].fill.fill_type is None


def test_monthly_audit_records_do_not_enter_daily_anomaly_workbook(tmp_path: Path) -> None:
    source = tmp_path / "summary-mismatch.xlsx"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    sheet["V5"] = 7
    sheet["W5"] = 70
    workbook.save(source)

    prepared = prepare_analysis(source, tmp_path / "work")
    finalized = finalize_analysis(prepared["analysis_file"], tmp_path / "outputs")

    assert finalized["summary_cells_corrected"] == 2
    anomaly = load_workbook(finalized["anomaly_workbook"], read_only=True)
    values = list(anomaly["异常审查"].iter_rows(values_only=True))
    assert len(values) == 2
    combined = "\n".join(str(value or "") for value in values[1])
    assert "月度汇总不一致" not in combined
    assert "出勤天数已对齐" not in combined


def test_summary_first_screening_avoids_parsing_every_daily_cell(tmp_path: Path) -> None:
    source = tmp_path / "large.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "月度汇总"
    sheet["A1"] = "月度汇总 统计日期：2026-07-01 至 2026-07-03"
    sheet["A2"] = "报表生成时间：2026-07-04 10:00"
    for column, header in {
        20: "上班缺卡次数",
        21: "下班缺卡次数",
        22: "迟到次数",
        23: "迟到时长",
        24: "早退次数",
        25: "早退时长",
        26: "旷工天数",
    }.items():
        sheet.cell(row=3, column=column).value = header
    sheet["AG3"] = "考勤结果"
    sheet["R3"] = "出勤天数"
    for offset in range(3):
        sheet.cell(row=4, column=33 + offset).value = str(offset + 1)
    for person_offset in range(20):
        row = 5 + person_offset
        sheet.cell(row=row, column=1).value = f"Person {person_offset}"
        for column in range(20, 27):
            sheet.cell(row=row, column=column).value = 0
        sheet.cell(row=row, column=18).value = 3
        for column in range(33, 36):
            sheet.cell(row=row, column=column).value = "研发A班:正常\n(09:00,18:00)"
    sheet["V24"] = 1
    sheet["W24"] = 10
    sheet["AH24"] = "研发A班:上班迟到10分钟\n(09:20,18:25)"
    workbook.save(source)

    prepared = prepare_analysis(source, tmp_path / "work")

    assert prepared["processed_people"] == 20
    assert prepared["daily_cells_total"] == 60
    assert prepared["summary_candidate_people"] == 1
    assert prepared["candidate_people"] == 1
    assert prepared["candidate_cells_examined"] == 3


def test_colored_authorized_leave_is_screened_but_not_reported(tmp_path: Path) -> None:
    source = tmp_path / "leave.xlsx"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    for column in range(20, 27):
        sheet.cell(row=5, column=column).value = 0
    sheet["AG5"].fill = PatternFill("solid", fgColor="FFFFFF")
    sheet["AH5"] = "研发A班:正常,病假07-02 09:00到07-02 18:00 1天\n(-,-)"
    sheet["AH5"].fill = PatternFill("solid", fgColor="D9EAD3")
    sheet["AI5"] = "研发A班:正常\n(09:00,18:00)"
    workbook.save(source)

    prepared = prepare_analysis(source, tmp_path / "work")

    assert prepared["summary_candidate_people"] == 0
    assert prepared["candidate_people"] == 1
    assert prepared["candidate_cells_examined"] == 1
    assert prepared["notable_cells"] == 0
    assert prepared["agent_review_items"] == 0


def test_white_unscheduled_punch_is_caught_by_lightweight_text_screen(tmp_path: Path) -> None:
    source = tmp_path / "unscheduled.xlsx"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    for column in range(20, 27):
        sheet.cell(row=5, column=column).value = 0
    for column in range(33, 36):
        sheet.cell(row=5, column=column).fill = PatternFill("solid", fgColor="FFFFFF")
    sheet["AG5"] = "研发A班:正常\n(09:00,18:00)"
    sheet["AH5"] = "未排班\n(09:00,18:35)"
    sheet["AI5"] = "研发A班:正常\n(09:00,18:00)"
    workbook.save(source)

    work_dir = tmp_path / "work"
    prepared = prepare_analysis(source, work_dir)

    assert prepared["summary_candidate_people"] == 0
    assert prepared["candidate_cells_examined"] == 1
    assert prepared["notable_cells"] == 1
    finalized = finalize_analysis(prepared["analysis_file"], tmp_path / "outputs")
    corrected = load_workbook(finalized["corrected_workbook"])
    assert corrected["月度汇总"]["AH5"].value == "未排班并打卡\n(09:00,18:35)"


def test_attendance_days_deducts_single_missing_and_rounds_monthly_minutes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "alignment.xlsx"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    sheet["R5"] = 3
    sheet["AG5"] = "研发A班:下班缺卡\n(09:00,-)"
    sheet["AH5"] = "研发A班:正常\n(09:36,18:10)"
    sheet["AI5"] = "研发A班:正常\n(09:00,18:00)"
    sheet["U5"] = 1
    sheet["V5"] = 1
    sheet["W5"] = 26
    sheet["X5"] = 1
    sheet["Y5"] = 26
    workbook.save(source)

    prepared = prepare_analysis(source, tmp_path / "work")
    finalized = finalize_analysis(prepared["analysis_file"], tmp_path / "outputs")

    corrected = load_workbook(finalized["corrected_workbook"])
    corrected_sheet = corrected["月度汇总"]
    assert corrected_sheet["R5"].value == 2.3
    assert corrected_sheet["W5"].value == 26
    assert corrected_sheet["Y5"].value == 26
    assert "基准出勤天数：3" in corrected_sheet["R5"].comment.text
    assert "迟到早退扣减：0.2 天" in corrected_sheet["R5"].comment.text
    assert finalized["attendance_days_cells_corrected"] == 1

    prepared_again = prepare_analysis(finalized["corrected_workbook"], tmp_path / "work-again")
    finalized_again = finalize_analysis(prepared_again["analysis_file"], tmp_path / "outputs-again")
    corrected_again = load_workbook(finalized_again["corrected_workbook"])
    assert corrected_again["月度汇总"]["R5"].value == 2.3
    assert finalized_again["attendance_days_cells_corrected"] == 0


def test_recorded_absence_does_not_get_missing_punch_deduction(tmp_path: Path) -> None:
    source = tmp_path / "absence.xlsx"
    _build_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["月度汇总"]
    sheet["R5"] = 2
    sheet["AG5"] = "研发A班:旷工\n(-,-)"
    sheet["AH5"] = "研发A班:正常\n(09:00,18:00)"
    sheet["AI5"] = "研发A班:正常\n(09:00,18:00)"
    sheet["V5"] = 0
    sheet["W5"] = 0
    sheet["Z5"] = 1
    workbook.save(source)

    prepared = prepare_analysis(source, tmp_path / "work")
    finalized = finalize_analysis(prepared["analysis_file"], tmp_path / "outputs")

    corrected = load_workbook(finalized["corrected_workbook"])
    assert corrected["月度汇总"]["R5"].value == 2
    assert finalized["attendance_days_cells_corrected"] == 0
