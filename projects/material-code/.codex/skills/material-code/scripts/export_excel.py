# -*- coding: utf-8 -*-
"""将 material_codes.db 反向导出为只含第一工作表的 Excel。

用途：把当前 SQLite 物料库（含新生成的物料）按原 Excel 的表头结构导回 Excel，
便于人工核对或与 ERP 原表对齐。数据库未存储的列（图纸编号、保质期/批次/品牌管理、
编号申请人、供应商编码、内部项目编码、备注等）导出时留空。

仅允许由 workflow.py 调用。
"""
import argparse
import datetime
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from matlib import (connect, DEFAULT_DB, sheet_names, read_sheet, SELECT_ALL)
from project_paths import find_project_root
from workflow_guard import require_workflow_entry
import openpyxl

# 表头文字 -> DB 字段名（与 matlib.connect 建表列一致）
HEADER_FIELD = {
    '大类': 'col_b', '大类编码': 'col_c', '一类': 'col_d',
    '一级子类编码': 'col_e', '二级': 'col_f', '二级子类编码': 'col_g',
    '上级分类-用友加': 'col_h', '旧物料编号': 'col_i', '新物料编码': 'col_j',
    '采购部门': 'col_k', '产品名称': 'name', '产品规格': 'spec',
    '单位': 'unit', '用途': 'usage', '引用': 'col_a', '编码示例': 'col_j',
}

# 这些列数据库未存储，导出留空
_EMPTY_HEADERS = {
    '图纸编号', '是否要做保质期管理', '是否要批次号管理', '是否要做品牌管理',
    '编号申请人', '供应商物料编码', '内部项目专用编码', '备注',
}


def field_for(header_name):
    """把表头文字映射到 DB 字段；空表头还原为 A 列引用串(col_a)。"""
    if header_name == '':
        return 'col_a'          # 第一页 A 列(空表头)存放引用拼接串
    if header_name in _EMPTY_HEADERS:
        return None
    return HEADER_FIELD.get(header_name)


def build_row(rec, header):
    """按表头顺序把一条 DB 记录映射成 Excel 行值列表。"""
    d = dict(zip([c.strip() for c in SELECT_ALL.split(',')], rec))
    out = []
    for h in header:
        f = field_for(h)
        out.append(d.get(f, '') if f else '')
    return out


def main():
    require_workflow_entry()
    ap = argparse.ArgumentParser()
    ap.add_argument('--db', default=DEFAULT_DB)
    ap.add_argument('--template', default='',
                    help='原清单 Excel，只读取第一工作表表头以还原格式')
    ap.add_argument('--out', default='',
                    help='导出路径；留空则输出到 <工程根>/output/物料清单导出_时间戳.xlsx')
    args = ap.parse_args()

    if not args.template or not os.path.exists(args.template):
        if args.template:
            print(f"[警告] 模板 Excel 不存在: {args.template}，将仅导出 A-T 标准列。")
        main_hdr = ['', '大类', '大类编码', '一类', '一级子类编码', '二级',
                    '二级子类编码', '上级分类-用友加', '旧物料编号', '新物料编码',
                    '采购部门', '图纸编号', '产品名称', '产品规格', '单位',
                    '是否要做保质期管理', '是否要批次号管理', '是否要做品牌管理',
                    '用途', '编号申请人']
        main_sheet = '物料清单（需求料号在此表查收或填写）'
    else:
        names = sheet_names(args.template)
        if not names:
            raise ValueError('模板 Excel 不包含工作表')
        main_sheet = names[0]
        main_hdr = read_sheet(args.template, main_sheet)[0]

    con = connect(args.db)
    cols = SELECT_ALL
    main_rows = con.execute(
        f"SELECT {cols} FROM materials WHERE retired=0 ORDER BY col_j").fetchall()

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = main_sheet
    ws1.append(main_hdr)
    for r in main_rows:
        ws1.append(build_row(r, main_hdr))

    root = find_project_root()
    out = args.out or str(root / 'output' /
                          f"物料清单导出_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"[导出完成] {out}")
    print(f"  主表({main_sheet}): {len(main_rows)} 条")
    print("  后续工作表: 未读取、未导出")
    print("  说明: 数据库未存 图纸编号/保质期·批次·品牌管理/编号申请人/供应商编码 等列，"
          "这些列导出留空。")

    # ---- 导出后自检：回读文件核对行数是否与 DB 一致 ----
    try:
        vwb = openpyxl.load_workbook(out, read_only=True)
        v_main = vwb[main_sheet].max_row - 1 if main_sheet in vwb.sheetnames else -1
        vwb.close()
        errs = []
        if v_main != len(main_rows):
            errs.append(f'主表行数不符: 文件={v_main} DB={len(main_rows)}')
        if errs:
            print('[警告] 导出自检:')
            for e in errs:
                print('   -', e)
        else:
            print('[通过] 导出自检：文件行数与数据库一致。')
    except Exception as e:
        print(f'[警告] 导出自检无法回读文件: {e}')
    con.close()


if __name__ == '__main__':
    main()
