import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "skills" / "manage-hr-resumes" / "scripts"))
from hr_resume.workbook import export_ledger, read_ledger


SNAPSHOT = {
    "revision": 3,
    "candidates": [{"id": "c1", "version": 1, "name": "张三", "source": "内推", "notes": "", "resume": {"path": "candidates/c1/a.txt", "sha256": "abc", "text_status": "ready"}}],
    "jobs": [{"id": "j1", "version": 2, "department": "研发", "title": "后端", "month": "2026-10", "headcount": 1, "status": "招聘中", "description": "", "criteria": [{"id": "x", "label": "Python", "weight": 1, "required": True}]}],
    "applications": [{"id": "a1", "version": 4, "candidate_id": "c1", "job_id": "j1", "month": "2026-10", "status": "待筛选", "owner": "", "notes": "", "next_action_date": ""}],
    "reviews": [], "interviews": [],
}


class WorkbookTests(unittest.TestCase):
    def test_export_and_read_only_changed_fields(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            path = Path(folder) / "ledger.xlsx"
            info = export_ledger(SNAPSHOT, path)
            self.assertEqual(info["revision"], 3)
            self.assertEqual(read_ledger(path), [])
            wb = load_workbook(path)
            wb["投递台账"]["F2"] = "二面"
            wb["投递台账"]["G2"] = "李四"
            wb.save(path)
            changes = read_ledger(path)
            self.assertEqual(changes[0]["id"], "a1")
            self.assertEqual(changes[0]["status"], "二面")
            self.assertEqual(changes[0]["owner"], "李四")

    def test_formula_is_stored_as_text_and_dates_are_native(self):
        snapshot = dict(SNAPSHOT)
        snapshot["applications"] = [dict(SNAPSHOT["applications"][0], notes="=HYPERLINK(\"http://evil\")", next_action_date="2026-11-01")]
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            path = Path(folder) / "ledger.xlsx"
            export_ledger(snapshot, path)
            wb = load_workbook(path, data_only=False)
            sheet = wb["投递台账"]
            self.assertEqual(sheet["H2"].value, "=HYPERLINK(\"http://evil\")")
            self.assertEqual(sheet["H2"].data_type, "s")
            self.assertEqual(sheet["I2"].value.date().isoformat(), "2026-11-01")
            self.assertEqual(sheet["I2"].number_format, "yyyy-mm-dd")

    def test_hidden_baseline_and_display_fields_are_integrity_protected(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            path = Path(folder) / "ledger.xlsx"
            export_ledger(SNAPSHOT, path)
            wb = load_workbook(path)
            wb["_台账基线"]["F2"] = "录用"
            wb.save(path)
            with self.assertRaises(ValueError):
                read_ledger(path)
            path2 = Path(folder) / "ledger2.xlsx"
            export_ledger(SNAPSHOT, path2)
            wb = load_workbook(path2)
            wb["投递台账"]["J2"] = "篡改姓名"
            wb.save(path2)
            with self.assertRaises(ValueError):
                read_ledger(path2)

    def test_rejects_relationship_edits_and_invalid_status(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            path = Path(folder) / "ledger.xlsx"
            export_ledger(SNAPSHOT, path)
            wb = load_workbook(path)
            wb["投递台账"]["C2"] = "c2"
            wb.save(path)
            with self.assertRaises(ValueError):
                read_ledger(path)
            export_ledger(SNAPSHOT, Path(folder) / "second.xlsx")
            wb = load_workbook(Path(folder) / "second.xlsx")
            wb["投递台账"]["F2"] = "未知状态"
            wb.save(Path(folder) / "second.xlsx")
            with self.assertRaises(ValueError):
                read_ledger(Path(folder) / "second.xlsx")

    def test_refuses_overwrite(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            path = Path(folder) / "ledger.xlsx"
            export_ledger(SNAPSHOT, path)
            with self.assertRaises(ValueError):
                export_ledger(SNAPSHOT, path)


if __name__ == "__main__":
    unittest.main()
