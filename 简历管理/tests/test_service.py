import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/manage-hr-resumes/scripts'
sys.path.insert(0, str(SCRIPTS))
from hr_resume.service import Service
from hr_resume.store import Store


class ServiceTests(unittest.TestCase):
    def setUp(self):
        runtime = Path(__file__).resolve().parents[1] / '.runtime'
        runtime.mkdir(exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=runtime)
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.root = self.base / '共享招聘'
        self.s = Service(self.root, '行政甲')
        self.s.execute('init')
        self.resume = self.base / '合成简历.txt'
        self.resume.write_text('曾负责电机机械设计，熟练使用 SolidWorks。参与三年生产工艺改进。', encoding='utf-8')
        self.c = self.s.execute('candidate-add', dict(name='示例甲', resume_path=str(self.resume)))
        self.jdata = dict(department='研发部', title='机械工程师', month='2026-10', headcount=2,
                          description='机械设计、SolidWorks 和工艺改进。', criteria=[
                              dict(id='design', label='机械设计经验', weight=60, required=True),
                              dict(id='cad', label='SolidWorks 能力', weight=40, required=False)])
        self.j = self.s.execute('jd-put', self.jdata)
        self.a = self.s.execute('application-add', dict(candidate_id=self.c['id'], job_id=self.j['id'], owner='行政甲'))

    def review(self):
        q = self.s.execute('review-queue')['items'][0]
        return dict(application_id=q['application_id'], application_version=q['application_version'], job_version=q['job_version'],
                    resume_sha256=q['resume_sha256'], text_sha256=q['text_sha256'], reviewer='测试 Agent',
                    criteria=[dict(criterion_id='design', rating=3, evidence='曾负责电机机械设计', reason='有相关项目描述，职责待面试核实'),
                              dict(criterion_id='cad', rating=4, evidence='熟练使用 SolidWorks', reason='有明确软件能力描述')],
                    strengths=['相关设计经验'], risks=['项目细节未提供'], questions=['请举例说明设计职责'])

    def test_complete_workflow_with_ledger_roundtrip(self):
        review = self.s.execute('review-submit', self.review())
        self.assertEqual(review['score'], 85)
        self.assertEqual(review['coverage'], 100)
        self.assertEqual(self.s.execute('review-queue')['items'], [])
        ranking = self.s.execute('ranking')['groups'][0]['ranked'][0]
        self.assertEqual((ranking['rank'], ranking['score']), (1, 85))
        interview = self.s.execute('interview-add', dict(application_id=self.a['id'], evaluator='业务主管',
            summary='项目描述清楚，待核实独立设计能力', recommendation='继续人工讨论', interview_date='2026-10-29', attachment_path=str(self.resume)))
        self.assertTrue((self.root / interview['attachment']['path']).exists())
        ledger = self.base / '总台账.xlsx'
        self.s.execute('ledger-export', dict(path=str(ledger)))
        wb = load_workbook(ledger)
        wb['投递台账'].cell(2, 6, '二面')
        wb['投递台账'].cell(2, 7, '行政乙')
        wb['投递台账'].cell(2, 9, '2026-11-02')
        wb.save(ledger)
        result = self.s.execute('ledger-import', dict(path=str(ledger)))
        self.assertEqual(result['updated'], 1)
        self.assertEqual(result['records'][0]['next_action_date'], '2026-11-02')
        self.assertEqual(self.s.execute('snapshot')['applications'][0]['status'], '二面')
        self.assertTrue(self.s.execute('doctor')['healthy'])

    def test_evidence_and_complete_criteria_enforced(self):
        review = self.review()
        review['criteria'][0]['evidence'] = '有十年航空经验'
        with self.assertRaisesRegex(ValueError, '证据'):
            self.s.execute('review-submit', review)
        self.assertEqual(self.s.execute('snapshot')['reviews'], [])
        review = self.review()
        review['criteria'][1]['criterion_id'] = 'design'
        with self.assertRaisesRegex(ValueError, '重复'):
            self.s.execute('review-submit', review)

    def test_unknown_not_false_rejection(self):
        review = self.review()
        for criterion in review['criteria']:
            criterion.update(rating=None, evidence='', reason='未提及，待核实')
        result = self.s.execute('review-submit', review)
        self.assertIsNone(result['score'])
        self.assertEqual(result['coverage'], 0)
        self.assertEqual(self.s.execute('ranking')['groups'][0]['ranked'], [])
        self.assertEqual(self.s.execute('snapshot')['applications'][0]['status'], '待复核')

    def test_jd_resume_and_transcript_invalidate_reviews(self):
        self.s.execute('review-submit', self.review())
        update = dict(self.jdata, id=self.j['id'], expected_version=1, description='新增项目要求')
        self.s.execute('jd-put', update)
        self.assertTrue(self.s.execute('snapshot')['reviews'][0]['is_stale'])
        self.assertEqual(self.s.execute('ranking')['groups'][0]['ranked'], [])
        self.s.execute('review-submit', self.review())
        transcript = self.base / '核对文本.txt'
        transcript.write_text('曾负责电机机械设计，熟练使用 SolidWorks。补充项目内容。', encoding='utf-8')
        self.s.execute('candidate-update', dict(id=self.c['id'], expected_version=1, text_path=str(transcript)))
        self.assertTrue(all(r['is_stale'] for r in self.s.execute('snapshot')['reviews']))

    def test_conflicts_lock_and_failed_transaction_leave_state_unchanged(self):
        self.s.execute('application-update', dict(id=self.a['id'], expected_version=1, owner='行政乙'))
        with self.assertRaisesRegex(ValueError, '版本冲突'):
            self.s.execute('application-update', dict(id=self.a['id'], expected_version=1, owner='行政甲'))
        with self.s.store.lock():
            with self.assertRaisesRegex(ValueError, '写入'):
                self.s.execute('application-update', dict(id=self.a['id'], expected_version=2, owner='行政丙'))
        before = self.s.execute('snapshot')
        original = self.s.store.atomic
        def fail_catalog(path, content):
            if path.name == 'catalog.json':
                raise OSError('模拟断线，提交索引前失败')
            return original(path, content)
        with patch.object(self.s.store, 'atomic', side_effect=fail_catalog):
            with self.assertRaises(OSError):
                self.s.execute('application-update', dict(id=self.a['id'], expected_version=2, owner='行政丙'))
        self.assertEqual(self.s.execute('snapshot'), before)
        self.s.execute('application-update', dict(id=self.a['id'], expected_version=2, owner='行政丁'))
        self.assertEqual(self.s.execute('snapshot')['applications'][0]['owner'], '行政丁')

    def test_batch_ledger_conflict_has_no_partial_updates(self):
        other = self.s.execute('candidate-add', dict(name='示例乙', resume_path=str(self.resume)))
        a2 = self.s.execute('application-add', dict(candidate_id=other['id'], job_id=self.j['id']))
        ledger = self.base / '台账.xlsx'
        self.s.execute('ledger-export', dict(path=str(ledger)))
        wb = load_workbook(ledger)
        wb['投递台账'].cell(2, 7, '同时编辑甲')
        wb['投递台账'].cell(3, 7, '同时编辑乙')
        wb.save(ledger)
        self.s.execute('application-update', dict(id=a2['id'], expected_version=1, owner='抢先修改'))
        before = self.s.execute('snapshot')
        with self.assertRaisesRegex(ValueError, '版本冲突'):
            self.s.execute('ledger-import', dict(path=str(ledger)))
        self.assertEqual(before, self.s.execute('snapshot'))

    def test_tampered_record_detected_and_root_boundary(self):
        catalog = self.s.store.catalog()
        rel = catalog['records']['jobs'][self.j['id']]['path']
        (self.root / rel).write_text('{}')
        with self.assertRaisesRegex(ValueError, '校验失败'):
            self.s.execute('snapshot')
        with self.assertRaisesRegex(ValueError, '超出'):
            self.s.store.safe('../escape.json')

    def test_cli_json_and_missing_root_no_mutation(self):
        result = subprocess.run([sys.executable, str(SCRIPTS / 'hr.py'), '--root', str(self.root), '--actor', '测试', 'snapshot'], capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(json.loads(result.stdout)['ok'])
        absent = self.base / '不存在'
        with self.assertRaisesRegex(ValueError, '初始化'):
            Service(absent, '测试').execute('snapshot')
        self.assertFalse(absent.exists())

    def test_explicit_reassessment_and_ocr_integrity(self):
        initial = self.s.execute('review-submit', self.review())
        q = self.s.execute('review-queue', {'include_reviewed': True})['items'][0]
        reassessment = dict(application_id=self.a['id'], application_version=q['application_version'], job_version=q['job_version'],
                            resume_sha256=q['resume_sha256'], text_sha256=q['text_sha256'], reviewer='人工复核Agent',
                            criteria=initial['criteria'], strengths=[], risks=[], questions=[])
        with self.assertRaisesRegex(ValueError, '已有有效审核'):
            self.s.execute('review-submit', reassessment)
        reassessment['supersedes_review_id'] = initial['id']
        reassessment['criteria'][0]['rating'] = 2
        new = self.s.execute('review-submit', reassessment)
        self.assertEqual(self.s.execute('ranking')['groups'][0]['ranked'][0]['review_id'], new['id'])
        ledger = self.base / '重审台账.xlsx'
        self.s.execute('ledger-export', {'path': str(ledger)})
        self.assertEqual(load_workbook(ledger)['岗位排序'].max_row, 2)
        transcript = self.base / 'OCR.txt'
        transcript.write_text('曾负责电机机械设计，熟练使用 SolidWorks', encoding='utf-8')
        updated = self.s.execute('candidate-update', dict(id=self.c['id'], expected_version=1, text_path=str(transcript)))
        self.s.execute('review-submit', self.review())
        (self.root / updated['resume']['transcript']['path']).write_text('篡改转录件', encoding='utf-8')
        self.assertFalse(self.s.execute('doctor')['healthy'])
        self.assertEqual(self.s.execute('ranking')['groups'][0]['ranked'], [])
        self.assertTrue(self.s.execute('review-queue')['blocked'])


if __name__ == '__main__':
    unittest.main()
