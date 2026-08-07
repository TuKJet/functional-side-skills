from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).parents[1]
SOURCE_SKILL = PROJECT_ROOT / ".codex/skills/material-code"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class PortableIntegrationTests(unittest.TestCase):
    def run_command(
        self, command: list[str], cwd: Path, expected: int = 0
    ) -> str:
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"
        proc = subprocess.run(
            command,
            cwd=cwd,
            env=env,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            check=False,
        )
        self.assertEqual(proc.returncode, expected, proc.stdout)
        return proc.stdout

    def test_sync_and_prepare_without_workbuddy(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            skill = root / ".codex/skills/material-code"
            skill.parent.mkdir(parents=True)
            shutil.copytree(SOURCE_SKILL, skill)
            (root / "pyproject.toml").write_text(
                '[project]\nname = "wuliao-bianma"\nversion = "0.0.0"\n',
                encoding="utf-8",
            )
            self.assertFalse((root / ".workbuddy").exists())

            headers = [
                "", "大类", "大类编码", "一类", "一级子类编码", "二级",
                "二级子类编码", "上级分类-用友加", "旧物料编号", "新物料编码",
                "采购部门", "图纸编号", "产品名称", "产品规格", "单位",
                "是否要做保质期管理", "是否要批次号管理", "是否要做品牌管理",
                "用途", "编号申请人",
            ]
            master = root / "master.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "正式第一页"
            first.append(headers)
            first.append(
                [
                    "原材料电子元器件线材", "原材料", "01", "电子元器件", "09",
                    "线材", "04", "010904", "", "0109040001", "供应链", "",
                    "测试射频线", "RF-1", "根", "", "", "", "测试用途", "",
                ]
            )
            second = workbook.create_sheet("禁止导入的第二页")
            second.append(headers)
            second.append(
                [
                    "错误数据", "错误类", "99", "错误一类", "99", "错误二类",
                    "99", "999999", "", "9999999999", "错误部门", "",
                    "第二页物料", "X", "个", "", "", "", "", "",
                ]
            )
            workbook.save(master)

            application = root / "application.xlsx"
            app_workbook = Workbook()
            app_sheet = app_workbook.active
            app_sheet.append(
                [
                    "产品名称", "产品规格", "单位", "用途", "编号申请人",
                    "是否要做保质期管理", "是否要批次号管理", "是否要做品牌管理",
                ]
            )
            app_sheet.append(
                ["测试射频线", "RF-1", "条", "测试用途", "申请人", "", "", ""]
            )
            app_workbook.save(application)

            workflow = skill / "scripts/workflow.py"
            sync_output = self.run_command(
                [sys.executable, str(workflow), "sync", str(master), "--confirmed"],
                root,
            )
            self.assertIn("后续工作表: 未读取、未导入、未参与流水号计算", sync_output)

            database = root / "material_codes.db"
            con = sqlite3.connect(database)
            rows = con.execute(
                "SELECT col_j,name FROM materials ORDER BY col_j"
            ).fetchall()
            con.close()
            self.assertEqual(rows, [("0109040001", "测试射频线")])

            before = sha256(database)
            prepare_output = self.run_command(
                [sys.executable, str(workflow), "prepare", str(application)], root
            )
            self.assertIn("[完全一致·复用]", prepare_output)
            self.assertEqual(before, sha256(database))

            active = json.loads(
                (root / "output/material-code-jobs/active.json").read_text(
                    encoding="utf-8"
                )
            )
            manifest = json.loads(
                Path(active["manifest"]).read_text(encoding="utf-8")
            )
            result = Path(manifest["result"])
            result_workbook = load_workbook(result, data_only=True)
            result_rows = list(result_workbook.active.iter_rows(values_only=True))
            self.assertEqual(result_rows[1][9], "0109040001")
            self.assertEqual(result_rows[1][14], "条")
            exact_reuse_row = list(result_workbook.active.iter_rows(min_row=2, max_row=2))[0]
            self.assertTrue(all(
                cell.fill.fill_type == "solid"
                and cell.fill.fgColor.rgb == "FFFFC7CE"
                and cell.font.color is not None
                and cell.font.color.type == "rgb"
                and cell.font.color.rgb == "FF9C0006"
                for cell in exact_reuse_row
            ))
            result_workbook.close()

            second_prepare = self.run_command(
                [sys.executable, str(workflow), "prepare", str(application)], root, 2
            )
            self.assertIn("该确认已被上一份申请使用", second_prepare)

            new_application = root / "new-application.xlsx"
            new_workbook = Workbook()
            new_sheet = new_workbook.active
            new_sheet.append(
                [
                    "产品名称", "产品规格", "单位", "用途", "编号申请人",
                    "是否要做保质期管理", "是否要批次号管理", "是否要做品牌管理",
                ]
            )
            new_sheet.append(
                ["新测试射频线", "RF-2", "条", "测试用途", "申请人", "", "", ""]
            )
            new_workbook.save(new_application)

            self.run_command(
                [sys.executable, str(workflow), "use-existing", "--confirmed"], root
            )
            self.run_command(
                [sys.executable, str(workflow), "prepare", str(new_application)], root
            )
            before_commit = sha256(database)
            commit_output = self.run_command(
                [sys.executable, str(workflow), "commit", "1:1"], root
            )
            self.assertIn("已生成预定编码", commit_output)
            self.assertEqual(before_commit, sha256(database))

            active = json.loads(
                (root / "output/material-code-jobs/active.json").read_text(
                    encoding="utf-8"
                )
            )
            manifest = json.loads(
                Path(active["manifest"]).read_text(encoding="utf-8")
            )
            generated_workbook = load_workbook(
                manifest["result"], read_only=True, data_only=True
            )
            generated_rows = list(
                generated_workbook.active.iter_rows(values_only=True)
            )
            generated_workbook.close()
            self.assertEqual(generated_rows[1][9], "0109040002")

            direct_child = self.run_command(
                [
                    sys.executable,
                    str(skill / "scripts/sync_db.py"),
                    "--excel",
                    str(master),
                    "--db",
                    str(database),
                ],
                root,
                2,
            )
            self.assertIn("禁止直接执行", direct_child)


if __name__ == "__main__":
    unittest.main()
